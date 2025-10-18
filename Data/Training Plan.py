#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Created on Tue Oct  7 17:06:13 2025

@author: adonischeng
"""
from docx import Document
import pandas as pd
import nltk
nltk.download('punkt')
nltk.download('punkt_tab')
from nltk.tokenize import sent_tokenize
import re
import json
import os

doc = Document("Training Plan.docx")


full_text = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

sections = {}
current_section = None
record = False  
for line in full_text:
    # start only after the real "1.0 Introduction" appears
    if re.match(r"^1\.0\s+Introduction$", line):
        record = True

    if not record:
        continue  # skip preface, version history, table of contents

    # detect section headings
    if re.match(r"^\d+(\.\d+)*\s+", line):
        current_section = line
        sections[current_section] = []
    elif current_section:
        sections[current_section].append(line)



cleaned_data = []

for section, content in sections.items():
    text = " ".join(content)
    text = re.sub(r"<.*?>", "", text)  # remove placeholders
    text = re.sub(r"\s+", " ", text)  # normalize spaces



document = {
    "document_title": "Training Plan (Cleaned)",
    "source_filename": "Training Plan.docx",
    "sections": []
}

for section_title, content_list in sections.items():
    section_number = section_title.split(" ")[0]
    title = " ".join(section_title.split(" ")[1:])
    
    section_data = {
        "number": section_number,
        "title": title,
        "content": " ".join(content_list),
    }
    document["sections"].append(section_data)

# Save to JSON file
output_path = os.path.splitext("Training Plan.docx")[0] + ".cleaned.json"
with open(output_path, "w", encoding="utf-8") as f:
    json.dump(document, f, ensure_ascii=False, indent=2)